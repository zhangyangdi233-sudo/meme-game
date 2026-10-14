#!/usr/bin/env python3
"""Build editable Chinese review documents from the current chapter inventories.

Use the bundled document Python runtime. This builder preserves every `zh`
string and source identifier; it does not extract game copy or translate it.
Run the document artifact-operation marker before invoking it to author files,
then render the DOCX with the packaged renderer and inspect every page.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FONT = "Heiti SC"
DEFAULT_STEM = "失语症_地下室第一章_剩余中文修改稿"
TITLE = "失语症 地下室第一章剩余中文修改稿"
INSTRUCTIONS = (
    "本稿按当前带地下室的第一章整理，供继续修改中文。已确认的开场白门旁白另存一稿，不在这里重复。",
    "直接修改每个编号下方的中文正文，保留编号和位置说明。交回后，再同步翻译成自然的口语日语。",
    "如果正文含有 %s、%d、{名称}、[b] 等占位符或格式标记，请保留标记及其顺序，只改前后的中文。多行文本请留意原有换行。",
    "十字路口暂无独立对白，操作提示列在对应分节。共用界面文案的修改可能同时影响后续章节。",
)


def load_sections(paths: list[Path]) -> tuple[list[dict], list[dict]]:
    sections: list[dict] = []
    inputs: list[dict] = []
    seen_sections: set[str] = set()
    seen_entries: set[str] = set()
    for path in paths:
        raw = path.read_bytes()
        payload = json.loads(raw)
        if not isinstance(payload, dict) or not isinstance(payload.get("sections"), list):
            raise ValueError(f"Expected sections array in {path}")
        inputs.append({
            "path": path.name,
            "sha256": hashlib.sha256(raw).hexdigest(),
            "metadata": {key: value for key, value in payload.items() if key != "sections"},
            "source_sections": [
                {**{key: value for key, value in section.items() if key != "entries"},
                 "entry_count": len(section.get("entries", []))}
                for section in payload["sections"]
            ],
        })
        for section in payload["sections"]:
            section_id, title = section.get("id"), section.get("title")
            if not isinstance(section_id, str) or not section_id or section_id in seen_sections:
                raise ValueError(f"Missing or duplicate section ID: {section_id!r}")
            if not isinstance(title, str) or not title:
                raise ValueError(f"Missing section title: {section_id}")
            if not isinstance(section.get("entries"), list):
                raise ValueError(f"Expected entries array: {section_id}")
            seen_sections.add(section_id)
            for entry in section["entries"]:
                entry_id = entry.get("id")
                if not isinstance(entry_id, str) or not entry_id or entry_id in seen_entries:
                    raise ValueError(f"Missing or duplicate entry ID: {entry_id!r}")
                if not isinstance(entry.get("zh"), str):
                    raise ValueError(f"Chinese source must be a string: {entry_id}")
                if not isinstance(entry.get("context", ""), str):
                    raise ValueError(f"Context must be a string: {entry_id}")
                for field in ("source_file", "source_locator"):
                    if field not in entry:
                        raise ValueError(f"Missing {field}: {entry_id}")
                seen_entries.add(entry_id)
            sections.append(section)
    if not seen_entries:
        raise ValueError("No entries to author")
    return sections, inputs


def set_font(style, size: float, color: str = "000000", bold: bool = False) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), FONT)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attr), None)


def clear_borders(element) -> None:
    for border in element.xpath(".//w:pBdr"):
        border.getparent().remove(border)


def make_document() -> Document:
    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin = page.bottom_margin = Cm(1.8)
    page.left_margin = page.right_margin = Cm(2)
    page.footer_distance = Cm(0.8)
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            set_font(style, 11.5)
            clear_borders(style.element)
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 22), ("Subtitle", 13), ("Heading 1", 15)):
        style = doc.styles[name]
        set_font(style, size, bold=name == "Heading 1")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(8)
    specs = (
        ("ReviewInstructions", 10.5, "000000", False, 6, False),
        ("ReviewEntryId", 11, "000000", True, 2, True),
        ("ReviewContext", 9.5, "555555", False, 3, True),
        ("ReviewChinese", 11.5, "000000", False, 11, False),
    )
    for name, size, color, bold, after, keep in specs:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        set_font(style, size, color, bold)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = keep
        style.paragraph_format.keep_together = True
    footer = page.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(doc.styles["Footer"], 9)
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    return doc


def display_heading(value: str) -> str:
    # Section metadata is a label, not source copy. Keep document headings plain.
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s\u3400-\u9fff]", " ", value)).strip()


def display_context(value: str) -> str:
    """Shorten repeated editorial metadata only; never rewrite Chinese copy."""
    replacements = {
        "第一章 CRT 与解锁后的手机实际复用；此源文案也被其他章节共用，回填时会同时影响它们。": "",
        "写入 log_text；手机打开时 DeskLog 显示，CRT 中同一操作会更新该状态。": "手机与 CRT 的操作反馈。",
        "当前嵌套工程启动时仍进行摄像头授权；与根目录新版独立体验实现不同。": "启动时的摄像头授权。",
        "全游戏共用设置；第一章可用 Esc/F10 打开。": "设置界面，可按 Esc 或 F10 打开。",
        "实际摄像头服务返回的状态或故障提示；全游戏共用。": "摄像头状态或故障提示。",
        "全游戏共用的当前主菜单、语言选择或退出确认。": "主菜单、语言选择或退出确认。",
        "跨章节共用。": "",
        "全游戏共用。": "",
    }
    for original, shortened in replacements.items():
        value = value.replace(original, shortened)
    return value.strip(" ；")


def build(args: argparse.Namespace) -> dict:
    sections, inputs = load_sections([args.story, args.shared])
    entries = [entry for section in sections for entry in section["entries"]]
    for number, entry in enumerate(entries, start=1):
        entry["review_id"] = f"C1-{number:03d}"
    authored_sections = [section for section in sections if section["entries"]]
    doc = make_document()
    doc.add_paragraph(TITLE, style="Title")
    doc.add_paragraph(f"中文编辑稿 {args.date}", style="Subtitle")
    for instruction in INSTRUCTIONS:
        doc.add_paragraph(instruction, style="ReviewInstructions")
    count_line = f"本稿共 {len(authored_sections)} 个分节，{len(entries)} 条中文。"
    doc.add_paragraph(count_line, style="ReviewInstructions")
    plain = [TITLE, f"中文编辑稿 {args.date}", "", *INSTRUCTIONS, count_line, ""]
    expected_copy: list[str] = []
    for number, section in enumerate(authored_sections, start=1):
        heading = f"{number:02d} {display_heading(section['title'])}"
        doc.add_paragraph(heading, style="Heading 1")
        plain.extend([heading, ""])
        for entry in section["entries"]:
            doc.add_paragraph(entry["review_id"], style="ReviewEntryId")
            plain.append(entry["review_id"])
            context = display_context(entry.get("context", ""))
            if context:
                doc.add_paragraph("位置  " + context, style="ReviewContext")
                plain.append("位置  " + context)
            body = doc.add_paragraph(entry["zh"], style="ReviewChinese")
            # Let very long authored passages flow instead of moving most of a
            # page into a blank gap. Ordinary records stay intact when possible.
            if len(entry["zh"]) > 500 or entry["zh"].count("\n") > 10:
                body.paragraph_format.keep_together = False
            expected_copy.append(entry["zh"])
            plain.extend([entry["zh"], ""])
    clear_borders(doc.element)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "当前带地下室第一章的剩余中文文本修改稿"
    doc.core_properties.author = ""
    doc.core_properties.keywords = ""
    args.output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = args.output_dir / f"{args.stem}.docx"
    txt_path = args.output_dir / f"{args.stem}.txt"
    doc.save(docx_path)
    txt_path.write_text("\n".join(plain).rstrip("\n") + "\n", encoding="utf-8")
    # The source map stays separate from the review body and preserves the
    # original inventory fields for a later Chinese/Japanese reintegration.
    map_path = args.mapping or args.output_dir / "回填资料" / "地下室第一章剩余中文定位.json"
    map_path.parent.mkdir(parents=True, exist_ok=True)
    map_path.write_text(json.dumps({
        "title": TITLE,
        "date": args.date,
        "entry_count": len(entries),
        "authored_section_count": len(authored_sections),
        "source_section_count": len(sections),
        "inputs": inputs,
        "sections": sections,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    loaded = Document(docx_path)
    actual_copy = [p.text for p in loaded.paragraphs if p.style.name == "ReviewChinese"]
    if actual_copy != expected_copy:
        raise AssertionError("DOCX round-trip did not preserve the Chinese source strings")
    with ZipFile(docx_path) as archive:
        for member in ("word/document.xml", "word/styles.xml"):
            if b"<w:pBdr" in archive.read(member):
                raise AssertionError(f"Unexpected paragraph border in {member}")
    actual_ids = [p.text for p in loaded.paragraphs if p.style.name == "ReviewEntryId"]
    if actual_ids != [entry["review_id"] for entry in entries]:
        raise AssertionError("DOCX review identifiers do not match the source map")
    return {"docx": str(docx_path), "txt": str(txt_path), "mapping": str(map_path),
            "sections": len(authored_sections), "entries": len(entries), "copy_verified": True}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--story", type=Path, required=True)
    parser.add_argument("--shared", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--mapping", type=Path)
    parser.add_argument("--stem", default=DEFAULT_STEM)
    parser.add_argument("--date", default="2026年10月13日")
    print(json.dumps(build(parser.parse_args()), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
