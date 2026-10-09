"""Build the concise, merged October 10 handoff from verified current behavior."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
out = root / 'docs/handoff/2026-10-10'
out.mkdir(parents=True, exist_ok=True)
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
for name, size in [('Normal', 11), ('Title', 23), ('Heading 1', 16), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != 'Normal'
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.paragraph_format.line_spacing = 1.10
    style.paragraph_format.space_after = Pt(5)
    if name != 'Normal': style.paragraph_format.keep_with_next = True
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)

def p(text, style=None):
    return doc.add_paragraph(text, style=style)

def h(text):
    p(text, 'Heading 2')

def page(title):
    para = p(title, 'Heading 1')
    para.paragraph_format.page_break_before = True

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + edge); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D8D8D8'); borders.append(b)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ['top', 'bottom', 'left', 'right']:
        m = OxmlElement('w:' + edge); m.set(qn('w:w'), '85'); m.set(qn('w:type'), 'dxa'); margins.append(m)
    pr.append(margins)
    for i, value in enumerate(headers): t.rows[0].cells[i].text = value
    header = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(header)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row): cells[i].text = value
    for row_index, row in enumerate(t.rows):
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            if row_index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8ECE7'); cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.08
                for run in para.runs:
                    run.font.size = Pt(10.5)
                    run.bold = row_index == 0
    gap = p('')
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = Pt(4)

p('Aphasia 修改与流程交接', 'Title')
p('合并简明版  2026年10月10日')
p('本说明合并上次交接与本轮修订。先按第1页运行和游玩，第2页查看改动与异常，第3页查制作入口。旧 Word、源模型和历史证据保留；出现冲突时，以本版现行流程为准。')
h('先启动同一个工程')
p('工程：D:/aphasia/meme-game/project.godot。使用 Godot 4.6.3；本机可运行 D:/aphasia/Start-Aphasia.ps1。开发验证用 D:/aphasia/Start-Chapter1-Dev.ps1，独立存档，不覆盖正常游玩进度。')
h('玩家怎样往下走')
p('开场全黑2秒 → 较暗白光门出现 → 向门走约8至9秒，中途响起敲门 → 听完、靠近按F → 慢开门且镜头向内推进 → 地下室教程五轮 → 空旷十字路口 → 远门 → 第二层。')
p('每轮：下楼，在CRT拾字、组句、投稿；退出屏幕后向本轮NPC交付。成功交付才可从左后方普通门离开。开门镜头连续推进，下一轮从楼梯顶进入；已取消旧12米黑隧道、循环白门和白闪。')
table(['轮次', '承接角色', '手机解锁'], [
    ['1', '护灯人', '信号瀑布'], ['2', '迟到者', '无新增应用'], ['3', '回声住户', '笔记本'], ['4', '抄写员', '无新增应用'], ['5', '无名信徒', '完成后通往十字路口'],
], [.7, 2, 4.35])
p('五轮都要交付；第5轮不再奖励巴别塔。十字路口远门检查五轮完成与两款App权限。已有字可反复用，CRT与手机共享同一词库、句子和规则。')
h('操作和保留规则')
p('WASD移动；F交互；Tab开关手机或退出CRT；Esc/F10设置；开发模式F9。Win/Alt-Tab失焦会放开鼠标，返回后点击游戏恢复视角。CRT进入和退出各有缓动镜头。')
p('资金、每日行动限制、巴别塔App、左侧资源栏、自动播放和历史入口已取消。污染继续在后台影响语言和后续楼层，玩家看不到数值。后续推进在有效交互完成时检查，不再等每日行动耗尽。')

page('这次改了什么')
table(['部分', '现在的行为'], [
    ['字词和界面', '可拾取字跟正文同字号；拖字松手后从落点继续掉落；滚轮加快；设置与CRT补齐中英日切换。'],
    ['镜头和门', '开场先黑2秒、拉远白门并减亮；门开时镜头持续推进。地下室用普通门接续；CRT交互距离缩小、进退缓动。'],
    ['空间和人物', '修补楼梯顶开口与穿插面；NPC仅水平朝向玩家，并按图片实际脚底贴地。'],
    ['陈设和材质', '桌子放大靠楼梯，木椅移到桌前且略拉开；小板凳隐藏；增箱子和桌面杂物；墙纸增加轻微旧感与局部痕迹，地毯加入织物法线。'],
    ['灯光和声音', '灯具缩小，灯泡发光且绑定微绿光源；CRT近处绿光加强。背景音乐降低5dB，增加四种地毯实录脚步。'],
    ['到达下一层', '世界画面保留，短暂显示大号无衬线楼层标题；不恢复旧每日结算界面。'],
], [1.18, 5.87])
h('异常清单 可直接作为以后彩蛋的入口')
table(['轮次', '异常', '玩家怎样处理'], [
    ['1', '正常灯光和单个墙钟', '建立对房间的记忆。'],
    ['2', '进屋后闪灯约3.6秒', '随后自行恢复；也可操作开关。'],
    ['3', '进屋后灯熄灭', '找到墙上开关按F，重新开灯。CRT绿光仍提供方向。'],
    ['4', 'CRT附近变成一整面钟', '18个钟，作为重复空间中的视觉异常。'],
    ['5', '墙钟缺半边', '缺失的手写数字只在摄像头手势X-ray窗口内出现。'],
    ['每轮', '出口小桌放花瓶或书', '每局有随机变化，同一存档保持一致。'],
], [.65, 2.25, 4.15])
p('异常集中在 basement_atmosphere.gd 顶部轮次常量和配置函数；改发生轮次或添加物件时，从这里开始。旧的三个X-ray出口箭头继续保留，正常画面不可见。')

page('接手制作从哪里开始')
table(['要改什么', '文件入口'], [
    ['状态与五轮流程', 'scripts/meme_game_state.gd\nscripts/progression/basement_loop_director.gd'],
    ['场景与门镜头', 'scripts/world/chapter_world.gd\n同目录 chapter_door_transition.gd'],
    ['灯 钟 家具 材质', 'scripts/world/basement_atmosphere.gd'],
    ['界面与屏幕', 'scripts/babel_meme_game.gd\nscripts/ui/chapter_terminal_session.gd'],
    ['字词物理和脚步', 'scripts/ui/word_physics_canvas.gd\nscripts/world/carpet_footsteps.gd'],
], [1.6, 5.45])
h('接任务 视频 和模型')
p('任务携带当前 npc_id、task_id、round_token，必须验证本轮投稿。旧稿与隐藏线索不替代交付；角色沿用旧贴图和对话，正式新正文待接。')
p('视频使用 ChapterVideoScreen.set_stream(stream) 与 play()；CRT独立4:3网格和0–1 UV保留。更换GLB时保留门轴、锚点、碰撞和X-ray标记，并同步 assets/chapter1/asset_contract.json。')
p('本轮陈设与修补由代码生成，原Blender保留在 source_assets/blender_basement_loop_v2。新贴图为项目生成；Mihacappy的CC0脚步来源与剪辑记录在 assets/audio/foley。')
h('旧存档与验证边界')
p('旧babel活动窗口迁回信号瀑布；旧资金与每日限制清理；旧隧道存档回安全入口。五轮任务、已获两App和隐藏物品仍分别校验，不能靠伪造单个权限跳关。')
p('本轮覆盖脚本、UI、物理、五轮CRT投稿交付与保存继续、门和场景渲染，执行清单见 verification.json。真实摄像头、用户视频、Windows系统键人工操作及耳机混音仍需实机体验；部分测试退出有RID/ObjectDB清理诊断。')
h('资料与版本')
p('最新交接：docs/handoff/2026-10-10；旧说明：相邻2026-10-09目录；原企划与文字稿：docs/design_originals、docs/narrative。研究出处：reference-research.md。')

path = out / 'Aphasia_修改与流程交接_合并简明版_20261010.docx'
doc.save(path)
print(path)
